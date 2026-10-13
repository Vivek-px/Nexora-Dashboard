"""Part 1B: live feeds, site profiles, benchmarking, bills & tariff, complaints, monthly PDF report and BRSR export."""
import calendar
import io
import uuid
from datetime import date
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field

import data
import images
import live
from engine.critical import sustainability_scorecard, waste_overflow_risk, maintenance_risk, cross_metric_events
from app import allowed, check_area, current_user, need, need_page, now
from db import DATA, one, q, run

router = APIRouter(prefix='/api')
MEASURED = data.MEASURED


def days_in(month):
    y, m = map(int, month.split('-'))
    return calendar.monthrange(y, m)[1]


def series_map(area, metric):
    return {m: (v, c) for m, v, c in data.merged(area, metric) if v is not None}


def evidence_of(classes):
    classes = set(classes)
    if classes == {MEASURED}:
        return 'Measured'
    if classes <= {'OBSERVED_OFFICIAL'}:
        return 'Official'
    if classes & {'SHARED_SYNTHETIC_BASELINE', 'MODELLED_AREA_SCENARIO'}:
        return 'Includes demonstration/modelled data'
    return 'Mixed'


# ------------------------------------------------------------------ site profiles

PROFILE_FIELDS = ['population', 'floor_area_m2', 'contract_demand_kw', 'energy_rate', 'demand_rate', 'min_billing_pct',
                  'excess_multiplier', 'fixed_charge', 'water_rate_per_kl', 'epi_target', 'lpcd_target', 'waste_target']


class Profile(BaseModel):
    population: float | None = Field(default=None, gt=0)
    floor_area_m2: float | None = Field(default=None, gt=0)
    contract_demand_kw: float | None = Field(default=None, gt=0)
    energy_rate: float | None = Field(default=None, ge=0)
    demand_rate: float | None = Field(default=None, ge=0)
    min_billing_pct: float | None = Field(default=80, ge=0, le=100)
    excess_multiplier: float | None = Field(default=2, ge=1, le=10)
    fixed_charge: float | None = Field(default=0, ge=0)
    water_rate_per_kl: float | None = Field(default=None, ge=0)
    epi_target: float | None = Field(default=None, gt=0)
    lpcd_target: float | None = Field(default=135, gt=0)
    waste_target: float | None = Field(default=None, gt=0)


def profile(area):
    p = one('SELECT * FROM site_profiles WHERE area=?', (area,)) or {}
    out = {k: p.get(k) for k in PROFILE_FIELDS}
    if out['population'] is None:  # the MVP's disclosed hypothetical population, until the site enters its own
        out['population'] = data.base()['water_population'].get(area)
        out['population_source'] = 'MVP hypothetical (1,000 users) – enter the real figure'
    else:
        out['population_source'] = 'Entered for this site'
    for k, default in (('min_billing_pct', 80), ('excess_multiplier', 2), ('fixed_charge', 0), ('lpcd_target', 135)):
        out[k] = default if out[k] is None else out[k]
    return out | {'area': area, 'updated_at': p.get('updated_at')}


@router.get('/profiles')
def get_profiles(user=Depends(current_user)):
    return {a: profile(a) for a in allowed(user)}


@router.put('/profiles/{area}')
def put_profile(area: str, p: Profile, user=Depends(need('admin'))):
    check_area(user, area)
    vals = [getattr(p, k) for k in PROFILE_FIELDS]
    run(f"""INSERT INTO site_profiles(area, {', '.join(PROFILE_FIELDS)}, updated_at) VALUES (?, {', '.join('?' * len(PROFILE_FIELDS))}, datetime('now'))
            ON CONFLICT(area) DO UPDATE SET {', '.join(f'{k}=excluded.{k}' for k in PROFILE_FIELDS)}, updated_at=excluded.updated_at""", (area, *vals))
    return profile(area)


# ------------------------------------------------------------------ benchmarking

def benchmarks(area):
    p = profile(area)
    pop, floor = p['population'], p['floor_area_m2']
    elec, water, waste, co2, peak = (series_map(area, k) for k in ('electricity_kwh', 'water_consumption_l', 'waste_kg', 'co2e_kg', 'peak_demand_kw'))
    months = sorted(elec)
    rows = []
    for m in months:
        d = days_in(m)
        e = elec.get(m, (None, None))[0]
        w = water.get(m, (None, None))[0]
        ws = waste.get(m, (None, None))[0]
        c = co2.get(m, (None, None))[0]
        pk = peak.get(m, (None, None))[0]
        rows.append({'month': m,
                     'kwh_per_person': e / pop if e is not None and pop else None,
                     'kwh_per_m2': e / floor if e is not None and floor else None,
                     'lpcd': w / (pop * d) if w is not None and pop else None,
                     'waste_kg_pp_day': ws / (pop * d) if ws is not None and pop else None,
                     'co2e_kg_per_person': c / pop if c is not None and pop else None,
                     'peak_vs_contract_pct': pk / p['contract_demand_kw'] * 100 if pk is not None and p['contract_demand_kw'] else None,
                     'evidence': evidence_of(x[m][1] for x in (elec, water, waste) if m in x)})
    last12 = months[-12:]
    epi = sum(elec[m][0] for m in last12) / floor if floor and len(last12) == 12 else None
    return {'area': area, 'profile': p, 'monthly': rows, 'epi_kwh_m2_year': epi, 'epi_period': f'{last12[0]} – {last12[-1]}' if last12 else '',
            'epi_evidence': evidence_of(elec[m][1] for m in last12) if last12 else None,
            'references': {'lpcd': {'value': p['lpcd_target'], 'label': 'CPHEEO norm 135 L/person/day (or the site target)'},
                           'epi': {'value': p['epi_target'], 'label': 'Site EPI target (kWh/m²/year) – set in Admin'},
                           'waste': {'value': p['waste_target'], 'label': 'Site waste target (kg/person/day) – set in Admin'}}}


@router.get('/benchmarks')
def get_benchmarks(user=Depends(need_page('benchmarks'))):
    return {a: benchmarks(a) for a in allowed(user)}


# ------------------------------------------------------------------ explainable critical intelligence

def _open_complaint_count(area, keywords):
    rows = q("SELECT category, title, description FROM complaints WHERE area=? AND status='open'", (area,))
    return sum(1 for row in rows if any(word in ' '.join(str(row.get(k) or '') for k in ('category', 'title', 'description')).lower() for word in keywords))


@router.get('/critical-insights')
def get_critical_insights(user=Depends(need_page('benchmarks'))):
    """Four explainable MVP capabilities; estimates are labelled and scoped to available data."""
    output = {}
    for area in allowed(user):
        b = benchmarks(area)
        rows = [dict(row) for row in b['monthly']]
        # Add annual energy intensity only when a complete trailing 12-month window and floor area exist.
        if rows:
            rows[-1]['epi_kwh_m2_year'] = b.get('epi_kwh_m2_year')
        card = sustainability_scorecard(rows, b['profile'])
        waste_points = data.merged(area, 'waste_kg')
        waste_complaints = _open_complaint_count(area, ('waste', 'overflow', 'bin', 'collection'))
        waste = waste_overflow_risk(waste_points, waste_complaints)
        series = {key: data.merged(area, key) for key in ('peak_demand_kw', 'hvac_kwh', 'leakage_probability_pct')}
        maint_complaints = _open_complaint_count(area, ('maintenance', 'equipment', 'hvac', 'pump', 'motor', 'electrical', 'leak'))
        maintenance = maintenance_risk(series, maint_complaints)
        screens = []
        for metric in ('electricity_kwh', 'water_consumption_l', 'waste_kg'):
            try:
                screens.append(data.screen_latest(area, metric))
            except (KeyError, IndexError, ValueError, TypeError):
                continue
        output[area] = {
            'area': area, 'area_name': data.area_name(area),
            'scorecard': card, 'waste_overflow': waste, 'maintenance_risk': maintenance,
            'cross_metric_events': cross_metric_events(screens),
            'data_period': rows[-1]['month'] if rows else None,
            'notice': 'Risk values are transparent rules-based screening estimates using available monthly/demo data, not calibrated probabilities. Distinguish synthetic/modelled inputs from measured facility uploads.'
        }
    return output


# ------------------------------------------------------------------ bills & tariff

def bill_for(p, kwh, peak_kw, water_l):
    if p['energy_rate'] is None:
        return None
    energy = kwh * p['energy_rate'] if kwh is not None else None
    cd = p['contract_demand_kw']
    demand = excess = None
    basis = None
    if cd and p['demand_rate'] is not None:
        # No peak reading for the month: the DISCOM still bills the minimum billing demand (a lower bound).
        billing = max(peak_kw, cd * p['min_billing_pct'] / 100) if peak_kw is not None else cd * p['min_billing_pct'] / 100
        basis = 'recorded peak' if peak_kw is not None else 'minimum billing demand (no peak reading)'
        demand = billing * p['demand_rate']
        excess = max(0.0, peak_kw - cd) * p['demand_rate'] * (p['excess_multiplier'] - 1) if peak_kw is not None else 0.0
    water = water_l / 1000 * p['water_rate_per_kl'] if water_l is not None and p['water_rate_per_kl'] is not None else None
    parts = [x for x in (energy, demand, excess, p['fixed_charge'] or 0) if x is not None]
    return {'energy': energy, 'demand': demand, 'excess_penalty': excess, 'fixed': p['fixed_charge'] or 0, 'demand_basis': basis,
            'electricity_total': sum(parts) if energy is not None else None, 'water': water}


def refit(series, horizon=3, upper=None):
    """Forecast the next months from a {month: (value, class)} map with the vendored MVP engine."""
    import pandas as pd
    from engine.forecasting import evaluate_forecast
    if not series:
        return {}
    s = pd.Series({pd.Timestamp(m + '-01'): v for m, (v, _) in series.items()}).sort_index()
    # Use the most recent contiguous run of months (the engine refuses gaps rather than interpolating).
    idx = pd.date_range(s.index.min(), s.index.max(), freq='MS').difference(s.index)
    if len(idx):
        s = s[s.index > idx.max()]
    res = evaluate_forecast(s, horizon, 0, upper)
    if not res.available:
        return {}
    num = lambda x: None if pd.isna(x) else float(x)
    return {'available': True, 'points': [[d.strftime('%Y-%m'), float(v), num(res.lower[d]), num(res.upper[d])] for d, v in res.points.items()]}


def bills(area):
    p = profile(area)
    elec, peak, water = series_map(area, 'electricity_kwh'), series_map(area, 'peak_demand_kw'), series_map(area, 'water_consumption_l')
    history = []
    for m in sorted(elec):
        b = bill_for(p, elec[m][0], peak.get(m, (None,))[0], water.get(m, (None,))[0])
        history.append({'month': m, 'kwh': elec[m][0], 'peak_kw': peak.get(m, (None,))[0], 'water_l': water.get(m, (None,))[0],
                        'evidence': evidence_of([elec[m][1]] + ([peak[m][1]] if m in peak else [])), 'bill': b})
    fc = data.base()['forecasts'][area]
    fe, fp = fc['electricity_kwh'].get('3', {}), fc['peak_demand_kw'].get('3', {})
    last_actual = history[-1]['month'] if history else ''
    # Measured uploads beyond the MVP export: re-run the MVP's own forecasting engine on the merged series.
    if fe.get('available') and not any(pt[0] > last_actual for pt in fe['points']):
        fe = refit(elec)
        fp = refit(peak, upper=None) if peak and max(peak) == last_actual else {}
    forecast = []
    if fe.get('available'):
        peaks = {pt[0]: pt[1] for pt in fp.get('points', [])} if fp.get('available') else {}
        for m, v, lo, hi in fe['points']:
            if m <= last_actual:
                continue  # measured uploads already cover this month
            forecast.append({'month': m, 'kwh': v, 'kwh_hi': hi, 'peak_kw': peaks.get(m), 'bill': bill_for(p, v, peaks.get(m), None),
                             'bill_hi': bill_for(p, hi, peaks.get(m), None) if hi is not None else None})
    recent = [h['bill']['electricity_total'] for h in history[-3:] if h['bill'] and h['bill']['electricity_total'] is not None]
    avg3 = sum(recent) / len(recent) if recent else None
    warnings = []
    for f in forecast:
        b = f['bill'] and f['bill']['electricity_total']
        if avg3 and b and b > avg3 * 1.10:
            warnings.append(f"Bill shock: {f['month']} electricity bill forecast ₹{b:,.0f} is {b / avg3 * 100 - 100:.0f}% above the last 3 months' average.")
        if p['contract_demand_kw'] and f['peak_kw'] and f['peak_kw'] > p['contract_demand_kw']:
            warnings.append(f"Contract demand: {f['month']} peak forecast {f['peak_kw']:.0f} kW exceeds the contract demand of {p['contract_demand_kw']:.0f} kW – excess-demand charges likely.")
    for h in history[-3:]:
        if p['contract_demand_kw'] and h['peak_kw'] and h['peak_kw'] > p['contract_demand_kw']:
            warnings.append(f"{h['month']}: recorded peak {h['peak_kw']:.0f} kW was above the contract demand ({p['contract_demand_kw']:.0f} kW).")
    return {'area': area, 'profile': p, 'configured': p['energy_rate'] is not None, 'history': history, 'forecast': forecast,
            'avg_last3': avg3, 'warnings': warnings,
            'note': 'Estimate from your tariff inputs – check against your DISCOM tariff order and actual bills. Months without a peak reading use the minimum billing demand, so the real demand charge may be higher. Forecast uses the MVP forecasting engine.'}


@router.get('/bills')
def get_bills(user=Depends(need_page('bills'))):
    return {a: bills(a) for a in allowed(user)}


# ------------------------------------------------------------------ live AQI & weather

@router.get('/live')
def get_live(refresh: bool = False, user=Depends(current_user)):
    out = {}
    for a in data.base()['areas']:
        if a['id'] not in allowed(user):
            continue
        ref = a['reference']
        out[a['id']] = {'air': live.air_for(a['id'], ref, refresh),
                        'weather': live.cached(f"weather:{a['id']}", lambda r=ref: live.fetch_weather(r['latitude'], r['longitude']), refresh)}
    return out


# ------------------------------------------------------------------ complaints

CATEGORIES = ['Water leak', 'No water / low pressure', 'Power failure', 'Electrical hazard', 'Waste overflow', 'Drainage / flooding',
              'Air quality / smoke', 'Building damage', 'Other']
SEVERITIES = ['Low', 'Medium', 'High']


@router.get('/complaints')
def complaints(user=Depends(current_user)):
    areas = allowed(user)
    return [c for c in q('SELECT c.*, u.name AS reporter FROM complaints c LEFT JOIN users u ON u.id=c.reporter_id ORDER BY c.id DESC') if c['area'] in areas]


@router.post('/complaints')
async def add_complaint(area: str = Form(...), category: str = Form(...), title: str = Form(..., min_length=3, max_length=150),
                        description: str = Form('', max_length=3000), severity: str = Form('Medium'), lat: float | None = Form(None),
                        lon: float | None = Form(None), place: str = Form('', max_length=300), loc_source: str = Form('', pattern='^(|map|gps|search|coordinates)$'),
                        photo: UploadFile | None = File(None), user=Depends(need_page('complaints'))):
    check_area(user, area)
    if category not in CATEGORIES or severity not in SEVERITIES:
        raise HTTPException(400, 'Unknown category or severity')
    if (lat is None) != (lon is None) or (lat is not None and not (-90 <= lat <= 90 and -180 <= lon <= 180)):
        raise HTTPException(400, 'Invalid map location')
    name = ''
    if photo is not None and photo.filename:
        try:
            name = images.store(await photo.read(), photo.filename)
        except images.BadFile as e:
            raise HTTPException(e.status, str(e))
    cid = run('INSERT INTO complaints(area, category, title, description, severity, lat, lon, photo, reporter_id, place, loc_source) VALUES (?,?,?,?,?,?,?,?,?,?,?)',
              (area, category, title, description, severity, lat, lon, name, user['id'], place.strip(), loc_source if lat is not None else ''))
    return one('SELECT c.*, u.name AS reporter FROM complaints c LEFT JOIN users u ON u.id=c.reporter_id WHERE c.id=?', (cid,))


class DismissIn(BaseModel):
    reason: str = Field(default='', max_length=500)
    action_id: str = Field(default='', max_length=64)  # the complaint's tracked action, so its status matches


def _complaint_for(cid, user):
    c = one('SELECT * FROM complaints WHERE id=?', (cid,))
    if not c:
        raise HTTPException(404, 'Complaint not found')
    check_area(user, c['area'])
    return c


def _action_status(action_id, user, area, title, status, text):
    if not action_id:
        return
    cur = one('SELECT status FROM action_state WHERE action_id=?', (action_id,))
    if not cur:
        run('INSERT INTO action_state(action_id, area, title, status, updated_by) VALUES (?,?,?,?,?)', (action_id, area, title, status, user['id']))
    elif cur['status'] != status:
        run('UPDATE action_state SET status=?, updated_at=datetime(\'now\'), updated_by=? WHERE action_id=?', (status, user['id'], action_id))
    run('INSERT INTO action_events(action_id, user_id, kind, text) VALUES (?,?,?,?)', (action_id, user['id'], 'status', text))


@router.post('/complaints/{cid}/dismiss')
def dismiss_complaint(cid: int, b: DismissIn, user=Depends(need('manager'))):
    """Dismissed complaints leave every list (Today, Recommendations, Indi, open complaints) and move to the Dismissed tab."""
    c = _complaint_for(cid, user)
    if c['status'] == 'dismissed':
        return c
    run("UPDATE complaints SET status='dismissed', dismissed_by=?, dismissed_at=?, dismiss_reason=? WHERE id=?",
        (user['id'], now(), b.reason.strip(), cid))
    _action_status(b.action_id, user, c['area'], f"Inspect and resolve: {c['title']}", 'Dismissed', 'Complaint dismissed' + (f': {b.reason.strip()}' if b.reason.strip() else ''))
    return one('SELECT c.*, u.name AS reporter FROM complaints c LEFT JOIN users u ON u.id=c.reporter_id WHERE c.id=?', (cid,))


@router.post('/complaints/{cid}/restore')
def restore_complaint(cid: int, b: DismissIn, user=Depends(need('manager'))):
    c = _complaint_for(cid, user)
    run("UPDATE complaints SET status='open', dismissed_by=NULL, dismissed_at=NULL, dismiss_reason='' WHERE id=?", (cid,))
    if c['status'] == 'dismissed':
        _action_status(b.action_id, user, c['area'], f"Inspect and resolve: {c['title']}", 'Open', 'Complaint restored')
    return one('SELECT c.*, u.name AS reporter FROM complaints c LEFT JOIN users u ON u.id=c.reporter_id WHERE c.id=?', (cid,))


@router.delete('/complaints/{cid}')
def delete_complaint(cid: int, action_id: str = Query(default='', max_length=64), user=Depends(need('admin'))):
    """Admins only: remove a complaint for good – the record, its photo, and its action history and proof files."""
    c = _complaint_for(cid, user)
    files = [c['photo']] if c['photo'] else []
    if action_id:
        files += [e['file'] for e in q("SELECT file FROM action_events WHERE action_id=? AND file!=''", (action_id,))]
        run('DELETE FROM action_events WHERE action_id=?', (action_id,))
        run('DELETE FROM action_state WHERE action_id=?', (action_id,))
    run('DELETE FROM complaints WHERE id=?', (cid,))
    for f in files:
        (DATA / 'uploads' / Path(f).name).unlink(missing_ok=True)
    return {'ok': True, 'deleted': cid, 'files_removed': len(files)}


@router.get('/complaints/meta')
def complaint_meta(_=Depends(current_user)):
    return {'categories': CATEGORIES, 'severities': SEVERITIES}


# ------------------------------------------------------------------ monthly PDF report

def _fmt(v, unit=''):
    if v is None:
        return '-'
    s = f'{v:,.0f}' if abs(v) >= 100 else f'{v:,.2f}'
    return f'{s} {unit}'.strip()


def monthly_pdf(area, month, user):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ink, mute, line = colors.HexColor('#17143B'), colors.HexColor('#5B5878'), colors.HexColor('#DCD8EE')
    H1 = ParagraphStyle('h1', fontName='Helvetica-Bold', fontSize=17, leading=21, textColor=ink, spaceAfter=2)
    H2 = ParagraphStyle('h2', fontName='Helvetica-Bold', fontSize=12, leading=15, textColor=ink, spaceBefore=10, spaceAfter=4)
    P = ParagraphStyle('p', fontName='Helvetica', fontSize=9, leading=12, textColor=ink)
    S = ParagraphStyle('s', fontName='Helvetica', fontSize=7.8, leading=10, textColor=mute)
    clean = lambda t: str(t).replace('₹', 'Rs ').replace('₂', '2').replace('³', '3').replace('µ', 'u').replace('≥', '>=').replace('–', '-').replace('—', '-').replace('→', '->')  # noqa: E731

    def table(rows, widths):
        t = Table([[Paragraph(clean(c), P) for c in r] for r in rows], colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#EEEAF9')), ('LINEBELOW', (0, 0), (-1, -1), 0.4, line),
                               ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3)]))
        return t

    b = data.base()
    y, mo = map(int, month.split('-'))
    prev = f'{y - 1 if mo == 1 else y}-{12 if mo == 1 else mo - 1:02d}'
    last_year = f'{y - 1}-{mo:02d}'
    story = [Paragraph(clean(f'Nexora monthly report · {data.area_name(area)}'), H1),
             Paragraph(clean(f'{calendar.month_name[mo]} {y} · generated {date.today().isoformat()} by {user["name"]}'), S), Spacer(1, 6)]
    rows = [['Domain', 'Indicator', 'This month', 'Previous month', 'Same month last year', 'Evidence']]
    labels = {MEASURED: 'Measured', 'OBSERVED_OFFICIAL': 'Official', 'SHARED_SYNTHETIC_BASELINE': 'Demonstration', 'MODELLED_AREA_SCENARIO': 'Modelled', 'MISSING': 'Missing'}
    for dmn in b['domains']:
        k = dmn['default']
        sm, unit = series_map(area, k), b['metrics'][k]['unit']
        cur = sm.get(month)
        rows.append([dmn['name'], b['metrics'][k]['name'], _fmt(cur[0], unit) if cur else '-', _fmt(sm[prev][0], unit) if prev in sm else '-',
                     _fmt(sm[last_year][0], unit) if last_year in sm else '-', labels.get(cur[1], cur[1]) if cur else 'No value'])
    story += [Paragraph('Key indicators', H2), table(rows, [26 * mm, 40 * mm, 28 * mm, 28 * mm, 30 * mm, 24 * mm])]

    al = q('SELECT * FROM alerts WHERE area=? AND month=? ORDER BY id', (area, month))
    story.append(Paragraph(f'Alerts raised for this month ({len(al)})', H2))
    story.append(table([['Level', 'Indicator', 'Status', 'Message']] + [[a['level'], b['metrics'].get(a['metric'], {}).get('name', a['metric']), a['status'], a['message']] for a in al],
                       [24 * mm, 32 * mm, 22 * mm, 98 * mm]) if al else Paragraph('No alerts for this month.', P))

    st = q('SELECT s.*, u.name AS owner FROM action_state s LEFT JOIN users u ON u.id=s.owner_id WHERE s.area=?', (area,))
    counts = {k: sum(1 for s in st if s['status'] == k) for k in ('Open', 'In progress', 'Done', 'Dismissed')}
    overdue = [s for s in st if s['due_date'] and s['due_date'] < date.today().isoformat() and s['status'] not in ('Done', 'Dismissed')]
    story.append(Paragraph('Action tracker', H2))
    story.append(Paragraph(clean(f"Tracked actions: {len(st)} · Open {counts['Open']} · In progress {counts['In progress']} · Done {counts['Done']} · Overdue {len(overdue)}"), P))
    if overdue:
        story.append(table([['Overdue action', 'Owner', 'Due']] + [[s['title'], s['owner'] or '-', s['due_date']] for s in overdue[:15]], [120 * mm, 30 * mm, 26 * mm]))

    cps = [c for c in q('SELECT * FROM complaints WHERE area=? ORDER BY id', (area,)) if c['created_at'][:7] == month]
    story.append(Paragraph(f'Complaints logged this month ({len(cps)})', H2))
    where = lambda c: (c.get('place') or '')[:70] or (f"{c['lat']:.5f}, {c['lon']:.5f}" if c['lat'] is not None else '-')
    story.append(table([['Category', 'Severity', 'Title', 'Location']] + [[c['category'], c['severity'], c['title'] + (' (dismissed)' if c.get('status') == 'dismissed' else ''), where(c)] for c in cps], [34 * mm, 18 * mm, 70 * mm, 54 * mm]) if cps else Paragraph('None.', P))

    sv = q('SELECT * FROM savings WHERE area=? ORDER BY id', (area,))
    story.append(Paragraph('Savings recorded (all time)', H2))
    if sv:
        story.append(table([['Indicator', 'Periods', 'Saved', 'Value', 'Evidence']] + [[b['metrics'][s['metric']]['name'], f"{s['baseline_from']}-{s['baseline_to']} vs {s['after_from']}-{s['after_to']}",
                                                                                     _fmt(s['units_saved'], s['unit']), f"Rs {s['rupees_saved']:,.0f}", s['evidence']] for s in sv],
                           [30 * mm, 50 * mm, 30 * mm, 24 * mm, 42 * mm]))
        story.append(Paragraph(clean(f"Total: Rs {sum(s['rupees_saved'] for s in sv):,.0f}"), P))
    else:
        story.append(Paragraph('No savings recorded yet.', P))

    bm = benchmarks(area)
    row = next((r for r in bm['monthly'] if r['month'] == month), None)
    story.append(Paragraph('Benchmarks', H2))
    if row:
        story.append(table(
            [['Measure', 'This month', 'Reference'],
             ['Electricity per person', _fmt(row['kwh_per_person'], 'kWh'), '-'],
             ['Water per person per day', _fmt(row['lpcd'], 'L'), _fmt(bm['references']['lpcd']['value'], 'L (CPHEEO / target)')],
             ['Waste per person per day', _fmt(row['waste_kg_pp_day'], 'kg'), _fmt(bm['references']['waste']['value'], 'kg (target)')],
             ['Energy Performance Index (12 months)', _fmt(bm['epi_kwh_m2_year'], 'kWh/m2/yr'), _fmt(bm['references']['epi']['value'], 'kWh/m2/yr (target)')]],
            [70 * mm, 50 * mm, 56 * mm]))
        story.append(Paragraph(clean(f"Population used: {_fmt(bm['profile']['population'])} ({bm['profile']['population_source']})."), S))
    bl = bills(area)
    hb = next((h for h in bl['history'] if h['month'] == month), None)
    story.append(Paragraph('Electricity bill estimate', H2))
    if bl['configured'] and hb and hb['bill']:
        x = hb['bill']
        story.append(Paragraph(clean(f"Energy Rs {_fmt(x['energy'])} + demand Rs {_fmt(x['demand'])} + excess-demand Rs {_fmt(x['excess_penalty'])} + fixed Rs {_fmt(x['fixed'])} = Rs {_fmt(x['electricity_total'])}"), P))
        for w in bl['warnings'][:4]:
            story.append(Paragraph(clean('Warning: ' + w), P))
    else:
        story.append(Paragraph('Tariff not set for this site (Admin -> Site profiles).', P))
    story += [Spacer(1, 10), Paragraph(clean('Evidence: "Measured" = uploaded facility readings; "Official" = CPCB/OSPCB; "Demonstration" = MVP demonstration '
                                             'baseline, not locality measurements. Screening flags are statistical review prompts, not hazard or compliance findings. '
                                             'Bill figures are estimates from the tariff entered in the dashboard.'), S)]
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=17 * mm, rightMargin=17 * mm, topMargin=15 * mm, bottomMargin=15 * mm,
                      title=f'Nexora {data.area_name(area)} {month}').build(story)
    return buf.getvalue()


@router.get('/reports/monthly.pdf')
def report_pdf(area: str, month: str = Query(pattern=r'^\d{4}-\d{2}$'), user=Depends(need_page('reports'))):
    check_area(user, area)
    pdf = monthly_pdf(area, month, user)
    return Response(pdf, media_type='application/pdf', headers={'Content-Disposition': f'attachment; filename="Nexora_{area}_{month}.pdf"'})


# ------------------------------------------------------------------ BRSR / ESG export

def fy_months(fy):
    return [f'{fy}-{m:02d}' for m in range(4, 13)] + [f'{fy + 1}-{m:02d}' for m in range(1, 4)]


@router.get('/reports/brsr.xlsx')
def brsr(fy: int = Query(ge=2000, le=2100), user=Depends(need_page('reports'))):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    months = fy_months(fy)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Principle 6 summary'
    ws.append([f'Nexora – BRSR Principle 6 (environment) data pack, FY {fy}-{str(fy + 1)[2:]} (April {fy} – March {fy + 1})'])
    ws.append(['Prepared from the dashboard. Check every figure and its evidence class before filing; demonstration data must not be reported.'])
    ws.append([])
    head = ['Site', 'Indicator (BRSR)', 'Value', 'Unit', 'Months with data', 'Evidence', 'Source indicator']
    ws.append(head)
    for c in ws[4]:
        c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='EEEAF9')
    detail = wb.create_sheet('Monthly detail')
    detail.append(['Site', 'Month', 'Electricity kWh', 'Water L', 'Waste kg', 'Plastic kg', 'Organic kg', 'Paper kg', 'CO2e kg', 'Evidence classes'])
    for area in allowed(user):
        maps = {k: series_map(area, k) for k in ('electricity_kwh', 'water_consumption_l', 'waste_kg', 'plastic_waste_kg', 'organic_waste_kg', 'paper_waste_kg', 'co2e_kg')}

        def total(k):
            pts = [maps[k][m] for m in months if m in maps[k]]
            return (sum(v for v, _ in pts), len(pts), evidence_of(c for _, c in pts)) if pts else (None, 0, 'No data')
        name = data.area_name(area)
        e, n, ev = total('electricity_kwh')
        ws.append([name, 'Total electricity consumption', round(e * 0.0036, 2) if e is not None else None, 'GJ', n, ev, 'electricity_kwh × 0.0036'])
        w, n, ev = total('water_consumption_l')
        ws.append([name, 'Total water consumption', round(w / 1000, 1) if w is not None else None, 'kL', n, ev, 'water_consumption_l ÷ 1000'])
        for k, label in (('plastic_waste_kg', 'Plastic waste'), ('organic_waste_kg', 'Other waste – organic'), ('paper_waste_kg', 'Other waste – paper'), ('waste_kg', 'Total waste generated')):
            v, n, ev = total(k)
            ws.append([name, label, round(v / 1000, 3) if v is not None else None, 'tonnes', n, ev, k])
        c, n, ev = total('co2e_kg')
        ws.append([name, 'Prototype CO2e estimate (not Scope 1/2 accounting)', round(c / 1000, 2) if c is not None else None, 't CO2e', n, ev, 'co2e_kg'])
        for m in months:
            vals = [maps[k].get(m, (None, None)) for k in maps]
            if any(v is not None for v, _ in vals):
                detail.append([name, m, *[v for v, _ in vals], ', '.join(sorted({c for _, c in vals if c}))])
    for sheet in (ws, detail):
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition': f'attachment; filename="Nexora_BRSR_FY{fy}-{str(fy + 1)[2:]}.xlsx"'})

# ------------------------------------------------------------------ local simulated-IoT demo and hourly baseline (synthetic data only)
from engine.hourly_demo import generate_readings, detect_hourly_events, forecast_next_24, recommendations

@router.get('/demo-intelligence')
def get_demo_intelligence(scenario: str = Query('normal', pattern='^(normal|leak|energy_spike|compound)$'), user=Depends(need_page('benchmarks'))):
    """Repeatable synthetic sensor scenario for evaluator demos; does not write to production data tables."""
    rows = generate_readings(72, scenario)
    events = detect_hourly_events(rows)
    forecast = forecast_next_24(rows)
    actions = recommendations(events, forecast)
    return {'data_label': 'SIMULATED — not live facility measurements', 'scenario': scenario,
            'readings': rows[-12:], 'events': events, 'forecast': forecast, 'recommendations': actions,
            'method_notes': ['Synthetic hourly data generated locally; no physical IoT hardware required.',
                             'Anomaly threshold is 50% above the baseline mean calculated from earlier readings.',
                             'Forecast is an hour-of-day historical-mean baseline, not a trained model.',
                             'Recommendation benefits are not quantified until verified with measured before/after readings.']}
