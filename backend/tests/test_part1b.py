"""Part 1B tests: AQI maths, live feeds (mocked), profiles, benchmarks, bills, complaints, PDF and BRSR exports."""
import io
import sys
from pathlib import Path

import httpx
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from test_backend import client, login, real_image  # noqa: E402,F401  (fixture reuse)


def test_aqi_matches_official_value():
    import live
    # Palasuni Dec-2025: OSPCB reported AQI 332 with PM2.5 161 µg/m³ – the CPCB formula gives the same index.
    assert live.sub_index('PM2.5', 161) == 332
    assert live.category(332) == 'Very Poor'
    assert live.sub_index('PM10', 280) == 230
    aqi, lead, subs, valid = live.compute_aqi({'PM2.5': 161, 'PM10': 280, 'NO2': 20})
    assert (aqi, lead, valid) == (332, 'PM2.5', True)
    assert live.compute_aqi({'NO2': 20})[3] is False  # not enough pollutants for a valid AQI


class FakeHTTP:
    def __init__(self, routes):
        self.routes, self.calls = routes, []

    def get(self, url, params=None):
        self.calls.append((url, params))
        for key, body in self.routes.items():
            if key in url:
                return httpx.Response(200, json=body, request=httpx.Request('GET', url))
        return httpx.Response(404, json={}, request=httpx.Request('GET', url))


WEATHER = {'current': {'temperature_2m': 31.2, 'relative_humidity_2m': 70, 'apparent_temperature': 36, 'precipitation': 0, 'wind_speed_10m': 9, 'weather_code': 2},
           'daily': {'time': ['2026-10-12', '2026-10-13', '2026-10-14'], 'temperature_2m_max': [33, 34, 32], 'temperature_2m_min': [25, 25, 24],
                     'precipitation_sum': [0, 70, 5], 'precipitation_probability_max': [10, 90, 40]}}
CPCB = {'records': [
    {'station': 'Palasuni, Bhubaneswar - OSPCB', 'last_update': '12-10-2026 13:00:00', 'pollutant_id': 'PM2.5', 'avg_value': '88'},
    {'station': 'Palasuni, Bhubaneswar - OSPCB', 'last_update': '12-10-2026 13:00:00', 'pollutant_id': 'PM10', 'avg_value': '140'},
    {'station': 'Palasuni, Bhubaneswar - OSPCB', 'last_update': '12-10-2026 13:00:00', 'pollutant_id': 'NO2', 'avg_value': '30'},
    {'station': 'IRC Village, Bhubaneswar - OSPCB', 'last_update': '12-10-2026 13:00:00', 'pollutant_id': 'PM2.5', 'pollutant_avg': '40'},
    {'station': 'IRC Village, Bhubaneswar - OSPCB', 'last_update': '12-10-2026 13:00:00', 'pollutant_id': 'PM10', 'pollutant_avg': '70'},
    {'station': 'IRC Village, Bhubaneswar - OSPCB', 'last_update': '12-10-2026 13:00:00', 'pollutant_id': 'OZONE', 'pollutant_avg': '20'}]}
AIR_MODEL = {'current': {'time': '2026-10-12T13:00', 'pm2_5': 55.0, 'pm10': 90.0, 'nitrogen_dioxide': 12.0, 'sulphur_dioxide': 5.0, 'ozone': 60.0, 'carbon_monoxide': 400.0}}


def test_live_official_then_modelled_fallback(client, monkeypatch):
    import live
    fake = FakeHTTP({'air-quality-api': AIR_MODEL, 'api.data.gov.in': CPCB, 'api.open-meteo.com': WEATHER})
    monkeypatch.setattr(live, '_client', lambda: fake)
    monkeypatch.setenv('DATA_GOV_IN_API_KEY', 'test-key')
    out = client.get('/api/live').json()
    p, k = out['palasuni']['air'], out['khandagiri']['air']
    assert p['kind'] == 'official' and p['station'].startswith('Palasuni') and p['aqi'] == live.sub_index('PM2.5', 88) == 193 and p['lead'] == 'PM2.5' and p['category'] == 'Moderate'
    assert out['irc_nayapalli']['air']['pollutants']['PM2.5'] == 40  # older 'pollutant_avg' field also understood
    assert k['kind'] == 'modelled' and 'not an official' in k['source']  # no Khandagiri CAAQMS station → labelled model
    assert out['palasuni']['weather']['daily'][1]['rain_mm'] == 70
    calls = len(fake.calls)
    client.get('/api/live')  # served from the 30-minute cache
    assert len(fake.calls) == calls
    monkeypatch.delenv('DATA_GOV_IN_API_KEY')
    m = client.get('/api/live', params={'refresh': True}).json()['palasuni']['air']
    assert m['kind'] == 'modelled' and m['pollutants']['CO'] == pytest.approx(0.4)


def test_live_offline_is_reported_not_crashing(client, monkeypatch):
    import live

    class Down:
        def get(self, *a, **k):
            raise httpx.ConnectError('offline')
    monkeypatch.setattr(live, '_client', lambda: Down())
    out = client.get('/api/live').json()['palasuni']
    assert out['air']['kind'] == 'unavailable' and 'error' in out['weather']


def test_profiles_benchmarks_and_bills(client):
    assert client.get('/api/bills').json()['palasuni']['configured'] is False
    r = client.put('/api/profiles/palasuni', json={'population': 1200, 'floor_area_m2': 8000, 'contract_demand_kw': 55, 'energy_rate': 7.5,
                                                   'demand_rate': 300, 'water_rate_per_kl': 20, 'epi_target': 40}).json()
    assert r['population'] == 1200 and r['population_source'] == 'Entered for this site'
    import data
    elec = {m: v for m, v, _ in data.merged('palasuni', 'electricity_kwh')}
    bm = client.get('/api/benchmarks').json()['palasuni']
    assert bm['epi_kwh_m2_year'] == pytest.approx(sum(elec[f'2025-{m:02d}'] for m in range(1, 13)) / 8000)
    dec = next(x for x in bm['monthly'] if x['month'] == '2025-12')
    water = {m: v for m, v, _ in data.merged('palasuni', 'water_consumption_l')}
    assert dec['lpcd'] == pytest.approx(water['2025-12'] / (1200 * 31))
    bills = client.get('/api/bills').json()['palasuni']
    h = next(x for x in bills['history'] if x['month'] == '2025-12')
    peak = h['peak_kw']
    assert h['bill']['energy'] == pytest.approx(elec['2025-12'] * 7.5)
    assert h['bill']['demand'] == pytest.approx(max(peak, 55 * 0.8) * 300)
    assert h['bill']['excess_penalty'] == pytest.approx(max(0, peak - 55) * 300)
    assert len(bills['forecast']) == 3 and any('Contract demand' in w for w in bills['warnings'])
    mgr_cannot = client.post('/api/users', json={'username': 'mgr2', 'name': 'M', 'role': 'manager', 'password': 'manager-pass'})
    assert login(client, 'mgr2', 'manager-pass').put('/api/profiles/palasuni', json={}).status_code == 403
    assert mgr_cannot.status_code == 200


def test_complaints_any_role_and_validation(client):
    client.post('/api/users', json={'username': 'staff', 'name': 'Staff', 'role': 'viewer', 'areas': 'khandagiri', 'password': 'viewer-pass'})
    staff = login(client, 'staff', 'viewer-pass')
    png = real_image('PNG')
    r = staff.post('/api/complaints', data={'area': 'khandagiri', 'category': 'Waste overflow', 'title': 'Bin overflowing at gate', 'severity': 'Medium', 'lat': '20.258', 'lon': '85.789'},
                   files={'photo': ('bin.png', io.BytesIO(png), 'image/png')})
    assert r.status_code == 200 and r.json()['photo'].endswith('.png')
    assert staff.post('/api/complaints', data={'area': 'palasuni', 'category': 'Waste overflow', 'title': 'Not my site'}).status_code == 403
    assert staff.post('/api/complaints', data={'area': 'khandagiri', 'category': 'Aliens', 'title': 'Bad category'}).status_code == 400
    assert staff.post('/api/complaints', data={'area': 'khandagiri', 'category': 'Other', 'title': 'Bad pin', 'lat': '20.2'}).status_code == 400
    assert [c['title'] for c in client.get('/api/complaints').json()] == ['Bin overflowing at gate']


def test_pdf_and_brsr(client):
    client.put('/api/profiles/palasuni', json={'energy_rate': 7.5, 'contract_demand_kw': 55, 'demand_rate': 300})
    r = client.get('/api/reports/monthly.pdf', params={'area': 'palasuni', 'month': '2025-12'})
    assert r.status_code == 200 and r.content.startswith(b'%PDF') and len(r.content) > 2000
    from pypdf import PdfReader
    text = ''.join(p.extract_text() for p in PdfReader(io.BytesIO(r.content)).pages)
    assert 'Palasuni' in text and 'Key indicators' in text and 'Rs' in text
    r = client.get('/api/reports/brsr.xlsx', params={'fy': 2024})
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    rows = list(wb['Principle 6 summary'].iter_rows(values_only=True))
    elec = next(x for x in rows if x[0] == 'Palasuni' and x[1] == 'Total electricity consumption')
    import data
    pts = {m: v for m, v, _ in data.merged('palasuni', 'electricity_kwh')}
    fy = [f'2024-{m:02d}' for m in range(4, 13)] + [f'2025-{m:02d}' for m in range(1, 4)]
    assert elec[2] == pytest.approx(round(sum(pts[m] for m in fy) * 0.0036, 2)) and elec[4] == 12 and 'demonstration' in elec[5]


def test_bill_forecast_moves_past_uploaded_months(client):
    client.put('/api/profiles/palasuni', json={'energy_rate': 7.5, 'contract_demand_kw': 55, 'demand_rate': 300})
    before = client.get('/api/bills').json()['palasuni']['forecast']
    body = 'area,metric,month,value\n' + ''.join(f'palasuni,electricity_kwh,2026-{m:02d},{28000 + m * 100}\n' for m in (1, 2, 3))
    assert client.post('/api/upload', files={'file': ('r.csv', io.BytesIO(body.encode()), 'text/csv')}).status_code == 200
    after = client.get('/api/bills').json()['palasuni']
    assert [f['month'] for f in before] == ['2026-01', '2026-02', '2026-03']  # MVP export forecast, unchanged
    assert [f['month'] for f in after['forecast']] == ['2026-04', '2026-05', '2026-06']  # re-run on measured data
    assert all(f['bill']['electricity_total'] > 0 for f in after['forecast'])
